from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from PIL import Image
from pypdf import PdfReader

from omniconvert.core.formats import FormatId
from omniconvert.core.registry import ConversionRegistry


def _make_pdf(tmp_path: Path) -> Path:
    from reportlab.pdfgen import canvas

    path = tmp_path / "two_page.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(72, 720, "Page One Hello")
    c.showPage()
    c.drawString(72, 720, "Page Two Hello")
    c.showPage()
    c.save()
    return path


def test_pdf_to_png_and_zip(registry: ConversionRegistry, tmp_path: Path) -> None:
    pdf = _make_pdf(tmp_path)
    png = tmp_path / "p1.png"
    registry.convert(
        operation_id="pdf.to_png",
        source=pdf,
        destination=png,
        settings={},
    )
    with Image.open(png) as img:
        assert img.format == "PNG"
        assert img.size[0] > 100

    zpath = tmp_path / "pages.zip"
    registry.convert(
        operation_id="pdf.pages_to_zip",
        source=pdf,
        destination=zpath,
        settings={},
    )
    with zipfile.ZipFile(zpath) as zf:
        names = zf.namelist()
        assert "page_0001.png" in names
        assert "page_0002.png" in names


def test_pdf_to_jpeg(registry: ConversionRegistry, tmp_path: Path) -> None:
    pdf = _make_pdf(tmp_path)
    jpg = tmp_path / "p1.jpg"
    registry.convert(
        operation_id="pdf.to_jpeg",
        source=pdf,
        destination=jpg,
        settings={},
    )
    with Image.open(jpg) as img:
        assert img.format == "JPEG"


def test_docx_roundtrip_and_pdf(registry: ConversionRegistry, tmp_path: Path) -> None:
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Hello justconvertit")
    doc.add_paragraph("Second line")
    doc.save(docx_path)

    txt = tmp_path / "out.txt"
    registry.convert(
        operation_id="office.docx_to_txt",
        source=docx_path,
        destination=txt,
        settings={},
    )
    assert "Hello justconvertit" in txt.read_text(encoding="utf-8")

    pdf = tmp_path / "from_docx.pdf"
    registry.convert(
        operation_id="office.docx_to_pdf_text",
        source=docx_path,
        destination=pdf,
        settings={},
    )
    assert pdf.read_bytes()[:4] == b"%PDF"

    back = tmp_path / "back.docx"
    registry.convert(
        operation_id="office.txt_to_docx",
        source=txt,
        destination=back,
        settings={},
    )
    assert back.stat().st_size > 0


def test_pdf_ops_listed(registry: ConversionRegistry) -> None:
    ops = {o.operation_id for o in registry.available_operations_for(FormatId.PDF)}
    assert "pdf.extract_text" in ops
    assert "pdf.to_png" in ops
    assert "pdf.pages_to_zip" in ops
    assert "pdf.page1_pdf" not in ops
    assert "pdf.split_pages" not in ops

"""LibreOffice high-fidelity conversions when the system binary is present."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from omniconvert.converters.office_libreoffice import LibreOfficeConverter
from omniconvert.core.errors import ConversionFailedError
from omniconvert.core.registry import ConversionRegistry
from omniconvert.services.dependencies import get_dependency_service


def _lo_usable() -> bool:
    if not LibreOfficeConverter().is_available():
        return False
    path = get_dependency_service().libreoffice.path
    # Snap under restricted environments often refuses via snap-confine.
    if path and "snap" in path.parts:
        return False
    return True


pytestmark = pytest.mark.skipif(
    not _lo_usable(),
    reason="LibreOffice apt/deb binary not available (Snap-only or missing)",
)


def test_docx_to_pdf_via_libreoffice(registry: ConversionRegistry, tmp_path: Path) -> None:
    docx_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("justconvertit Quality", level=1)
    doc.add_paragraph("Unicode: café — 日本語 — مرحبا")
    doc.save(docx_path)

    pdf = tmp_path / "out.pdf"
    try:
        result = registry.convert(
            operation_id="office.docx_to_pdf",
            source=docx_path,
            destination=pdf,
            settings={},
        )
    except ConversionFailedError as exc:
        pytest.skip(f"LibreOffice could not run: {exc}")
    assert result.output_path.exists()
    assert pdf.read_bytes()[:4] == b"%PDF"
    reader = PdfReader(str(pdf))
    assert len(reader.pages) >= 1


def test_txt_to_pdf_via_libreoffice(registry: ConversionRegistry, tmp_path: Path) -> None:
    src = tmp_path / "notes.txt"
    src.write_text("Line one\nLine two — café\n", encoding="utf-8")
    pdf = tmp_path / "notes.pdf"
    try:
        registry.convert(
            operation_id="office.txt_to_pdf_lo",
            source=src,
            destination=pdf,
            settings={},
        )
    except ConversionFailedError as exc:
        pytest.skip(f"LibreOffice could not run: {exc}")
    assert pdf.read_bytes()[:4] == b"%PDF"

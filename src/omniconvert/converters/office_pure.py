"""Office conversions using free, pip-bundlable pure-Python libraries.

These work on Windows and Linux without installing LibreOffice.
Layout fidelity is limited (text-focused). LibreOffice remains an optional
high-fidelity upgrade when present on the system.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from omniconvert.core.converter import (
    CancelCheck,
    ConversionOperation,
    Converter,
    ProgressCallback,
    check_cancelled,
)
from omniconvert.core.errors import ConversionFailedError, ValidationError
from omniconvert.core.formats import FormatId
from omniconvert.core.job import ConversionResult
from omniconvert.core.limits import MAX_TEXT_BYTES
from omniconvert.core.output import AtomicWriter, verify_output


def _read_text_file(path: Path) -> str:
    if path.stat().st_size > MAX_TEXT_BYTES:
        raise ValidationError(
            f"Text file exceeds {MAX_TEXT_BYTES // (1024 * 1024)} MiB limit."
        )
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValidationError("Could not decode text file.")


def _text_to_pdf(text: str, destination: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    with AtomicWriter(destination) as tmp:
        c = canvas.Canvas(str(tmp), pagesize=letter)
        width, height = letter
        left, top, bottom = 50, height - 50, 50
        y = top
        c.setFont("Helvetica", 10)
        for line in text.splitlines() or [""]:
            while line:
                chunk = line[:95]
                line = line[95:]
                if y < bottom:
                    c.showPage()
                    c.setFont("Helvetica", 10)
                    y = top
                safe = chunk.encode("latin-1", errors="replace").decode("latin-1")
                c.drawString(left, y, safe)
                y -= 14
        c.save()


class OfficePureConverter(Converter):
    """Always available Office/text bridges (no external install)."""

    def operations(self) -> list[ConversionOperation]:
        return [
            ConversionOperation(
                operation_id="office.docx_to_txt",
                label="DOCX → TXT",
                source_formats=frozenset({FormatId.DOCX}),
                target_format=FormatId.TXT,
                description="Extract paragraph text from a Word document.",
                tested=True,
            ),
            ConversionOperation(
                operation_id="office.txt_to_docx",
                label="TXT → DOCX",
                source_formats=frozenset({FormatId.TXT}),
                target_format=FormatId.DOCX,
                description="Wrap plain text into a simple Word document.",
                tested=True,
            ),
            ConversionOperation(
                operation_id="office.docx_to_pdf_text",
                label="DOCX → PDF (text)",
                source_formats=frozenset({FormatId.DOCX}),
                target_format=FormatId.PDF,
                description=(
                    "Portable text-only PDF from DOCX (no LibreOffice). "
                    "Hidden when LibreOffice is installed — that path keeps real layout."
                ),
                wheel_label="PDF (text)",
                tested=True,
            ),
            ConversionOperation(
                operation_id="office.pptx_to_txt",
                label="PPTX → TXT",
                source_formats=frozenset({FormatId.PPTX}),
                target_format=FormatId.TXT,
                description="Extract text from PowerPoint slides.",
                tested=True,
            ),
            ConversionOperation(
                operation_id="office.pptx_to_pdf_text",
                label="PPTX → PDF (text)",
                source_formats=frozenset({FormatId.PPTX}),
                target_format=FormatId.PDF,
                description=(
                    "Portable text-only PDF from PPTX (no LibreOffice). "
                    "Hidden when LibreOffice is installed."
                ),
                wheel_label="PDF (text)",
                tested=True,
            ),
        ]

    def convert(
        self,
        *,
        operation_id: str,
        source: Path,
        destination: Path,
        settings: dict[str, Any],
        progress: ProgressCallback | None = None,
        is_cancelled: CancelCheck | None = None,
    ) -> ConversionResult:
        _ = settings
        check_cancelled(is_cancelled)
        if progress:
            progress(0.2, "Reading…")
        try:
            if operation_id == "office.docx_to_txt":
                text = self._docx_text(source)
                with AtomicWriter(destination) as tmp:
                    Path(tmp).write_text(text, encoding="utf-8")
                msg = "Extracted text from DOCX."
            elif operation_id == "office.txt_to_docx":
                from docx import Document

                text = _read_text_file(source)
                check_cancelled(is_cancelled)
                doc = Document()
                for para in text.splitlines() or [""]:
                    doc.add_paragraph(para)
                with AtomicWriter(destination) as tmp:
                    doc.save(str(tmp))
                msg = "Created DOCX from text."
            elif operation_id == "office.docx_to_pdf_text":
                text = self._docx_text(source)
                check_cancelled(is_cancelled)
                if progress:
                    progress(0.6, "Writing PDF…")
                _text_to_pdf(text, destination)
                msg = "Created text PDF from DOCX (layout simplified)."
            elif operation_id == "office.pptx_to_txt":
                text = self._pptx_text(source)
                with AtomicWriter(destination) as tmp:
                    Path(tmp).write_text(text, encoding="utf-8")
                msg = "Extracted text from PPTX."
            elif operation_id == "office.pptx_to_pdf_text":
                text = self._pptx_text(source)
                check_cancelled(is_cancelled)
                if progress:
                    progress(0.6, "Writing PDF…")
                _text_to_pdf(text, destination)
                msg = "Created text PDF from PPTX (layout simplified)."
            else:
                raise ConversionFailedError("Unknown office operation.")

            size = verify_output(destination, min_bytes=0)
            if progress:
                progress(1.0, "Done")
            return ConversionResult(destination, size, msg)
        except ValidationError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ConversionFailedError(
                "Office conversion failed.",
                detail=str(exc),
            ) from exc

    def _docx_text(self, source: Path) -> str:
        from docx import Document

        try:
            doc = Document(str(source))
        except Exception as exc:  # noqa: BLE001
            raise ValidationError("Could not read DOCX file.", detail=str(exc)) from exc
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also tables
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append("\t".join(cells))
        return "\n".join(parts) + ("\n" if parts else "")

    def _pptx_text(self, source: Path) -> str:
        from pptx import Presentation

        try:
            prs = Presentation(str(source))
        except Exception as exc:  # noqa: BLE001
            raise ValidationError("Could not read PPTX file.", detail=str(exc)) from exc
        parts: list[str] = []
        for i, slide in enumerate(prs.slides, start=1):
            parts.append(f"--- Slide {i} ---")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    parts.append(shape.text)
        return "\n".join(parts) + ("\n" if parts else "")
